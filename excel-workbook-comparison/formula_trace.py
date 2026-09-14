"""Parse an Excel formula, walk what it depends on, and re-evaluate it under scenarios.

The value comparison answers "this cell changed". To answer "and here is why, and what it did
to ROA", you have to read the *formula* and follow it down to the inputs that moved.

Three things live here:

  parse(text)                -> an AST for an Excel formula
  trace(book, sheet, cell)   -> the precedent tree, and the changed inputs inside it
  compile_scenario(...)      -> that tree flattened to "ROA as a function of the changed inputs",
                                so any subset of them can be switched old -> new and re-evaluated

That last part is the point. ROA is a ratio, so its drivers are *not* additive: the effect of two
changes together is not the sum of their separate effects. Re-evaluating the real formula under a
chosen subset is the only way to get that right.

Formulas are only available from .xlsx/.xlsm. pyxlsb and xlrd cannot see formula text, so a .xlsb
has to be saved as .xlsx before it can be traced.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

__all__ = [
    "FormulaError", "UnsupportedFormula", "parse", "evaluate", "Node", "remap_rows",
    "Book", "load_book", "Component", "StructuralChange", "Trace", "trace",
    "compile_scenario", "trace_columns", "eval_dag",
]


class FormulaError(ValueError):
    """The formula could not be parsed or evaluated."""


class UnsupportedFormula(FormulaError):
    """The formula uses something this evaluator deliberately does not implement."""


# ---------------------------------------------------------------------------
#  Tokenizer
# ---------------------------------------------------------------------------

_SHEET = r"(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!"
_A1 = r"\$?[A-Za-z]{1,3}\$?\d{1,7}"
_ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
           "#SPILL!", "#CALC!", "#GETTING_DATA")

_TOKEN_RE = re.compile(
    r"""(?P<ws>\s+)
      | (?P<error>\#(?:NULL!|DIV/0!|VALUE!|REF!|NAME\?|NUM!|N/A|SPILL!|CALC!|GETTING_DATA))
      | (?P<string>"(?:[^"]|"")*")
      | (?P<range>(?:""" + _SHEET + r""")?""" + _A1 + r"""\s*:\s*(?:""" + _SHEET + r""")?""" + _A1 + r""")
      | (?P<wholecol>(?:""" + _SHEET + r""")?\$?[A-Za-z]{1,3}\s*:\s*(?:""" + _SHEET + r""")?\$?[A-Za-z]{1,3}\b)
      | (?P<wholerow>(?:""" + _SHEET + r""")?\$?\d{1,7}\s*:\s*(?:""" + _SHEET + r""")?\$?\d{1,7}\b)
      | (?P<func>[A-Za-z_][\w.]*\s*(?=\())
      | (?P<ref>(?:""" + _SHEET + r""")?""" + _A1 + r""")
      | (?P<bool>\b(?:TRUE|FALSE)\b)
      | (?P<number>(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)
      | (?P<op><=|>=|<>|[-+*/^&=<>%])
      | (?P<lparen>\()
      | (?P<rparen>\))
      | (?P<comma>[,;])
    """, re.VERBOSE)


@dataclass
class Token:
    kind: str
    text: str
    pos: int


def tokenize(src: str) -> List[Token]:
    out, i, n = [], 0, len(src)
    while i < n:
        m = _TOKEN_RE.match(src, i)
        if not m or m.end() == i:
            raise FormulaError(f"cannot read formula at position {i}: {src[i:i + 24]!r}")
        kind = m.lastgroup
        if kind != "ws":
            out.append(Token(kind, m.group().strip(), i))
        i = m.end()
    return out


# ---------------------------------------------------------------------------
#  AST
# ---------------------------------------------------------------------------

@dataclass
class Node:
    kind: str                       # num | str | bool | err | ref | range | binop | unary | func
    value: Any = None               # literal value, operator symbol, or function name
    args: List["Node"] = field(default_factory=list)
    sheet: Optional[str] = None     # ref / range only
    coord: Any = None               # ref: (row, col); range: ((r1,c1),(r2,c2))


def _split_ref(text: str) -> Tuple[Optional[str], str]:
    if "!" in text:
        sheet, rest = text.rsplit("!", 1)
        sheet = sheet.strip()
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        return sheet, rest
    return None, text


def _a1(text: str) -> Tuple[int, int]:
    m = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?(\d{1,7})", text.strip())
    if not m:
        raise FormulaError(f"not a cell reference: {text!r}")
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


# ---------------------------------------------------------------------------
#  Parser - precedence climbing
# ---------------------------------------------------------------------------

_BINARY = [
    {"=", "<>", "<", ">", "<=", ">="},
    {"&"},
    {"+", "-"},
    {"*", "/"},
    {"^"},
]


class _Parser:
    def __init__(self, tokens: List[Token], src: str):
        self.t, self.i, self.src = tokens, 0, src

    def peek(self) -> Optional[Token]:
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self) -> Token:
        tok = self.peek()
        if tok is None:
            raise FormulaError(f"formula ends unexpectedly: {self.src!r}")
        self.i += 1
        return tok

    def expect(self, kind: str) -> Token:
        tok = self.take()
        if tok.kind != kind:
            raise FormulaError(f"expected {kind}, found {tok.text!r} in {self.src!r}")
        return tok

    def parse(self) -> Node:
        node = self.expr(0)
        if self.peek() is not None:
            raise FormulaError(f"unexpected {self.peek().text!r} in {self.src!r}")
        return node

    def expr(self, level: int) -> Node:
        if level >= len(_BINARY):
            return self.unary()
        node = self.expr(level + 1)
        while True:
            tok = self.peek()
            if tok is None or tok.kind != "op" or tok.text not in _BINARY[level]:
                return node
            self.take()
            rhs = self.expr(level + 1)
            node = Node("binop", tok.text, [node, rhs])

    def unary(self) -> Node:
        tok = self.peek()
        if tok is not None and tok.kind == "op" and tok.text in {"-", "+"}:
            self.take()
            return Node("unary", tok.text, [self.unary()])
        return self.postfix()

    def postfix(self) -> Node:
        node = self.primary()
        while True:
            tok = self.peek()
            if tok is not None and tok.kind == "op" and tok.text == "%":
                self.take()
                node = Node("unary", "%", [node])
            else:
                return node

    def primary(self) -> Node:
        tok = self.take()
        if tok.kind == "number":
            return Node("num", float(tok.text))
        if tok.kind == "string":
            return Node("str", tok.text[1:-1].replace('""', '"'))
        if tok.kind == "bool":
            return Node("bool", tok.text.upper() == "TRUE")
        if tok.kind == "error":
            return Node("err", tok.text)
        if tok.kind == "ref":
            sheet, a1 = _split_ref(tok.text)
            return Node("ref", tok.text, sheet=sheet, coord=_a1(a1))
        if tok.kind == "range":
            left, right = tok.text.split(":")
            s1, a1 = _split_ref(left)
            s2, a2 = _split_ref(right)
            return Node("range", tok.text, sheet=s1 or s2, coord=(_a1(a1), _a1(a2)))
        if tok.kind in ("wholecol", "wholerow"):
            return Node("wholeref", tok.text)
        if tok.kind == "func":
            name = tok.text.upper()
            self.expect("lparen")
            args: List[Node] = []
            if self.peek() is not None and self.peek().kind == "rparen":
                self.take()
                return Node("func", name, args)
            while True:
                args.append(self.expr(0))
                nxt = self.take()
                if nxt.kind == "rparen":
                    break
                if nxt.kind != "comma":
                    raise FormulaError(f"expected , or ) in {name}(...), found {nxt.text!r}")
            return Node("func", name, args)
        if tok.kind == "lparen":
            node = self.expr(0)
            self.expect("rparen")
            return node
        raise FormulaError(f"unexpected {tok.text!r} in {self.src!r}")


_PARSE_CACHE: Dict[str, Node] = {}


def parse(text: str) -> Node:
    """Parse an Excel formula. The leading '=' is optional.

    Cached: the same formula text recurs on every sheet and every month column, and the
    parsed tree is only ever read - the rebuilder allocates its own nodes.
    """
    hit = _PARSE_CACHE.get(text)
    if hit is not None:
        return hit
    src = text.strip()
    if src.startswith("="):
        src = src[1:]
    if not src:
        raise FormulaError("empty formula")
    node = _Parser(tokenize(src), src).parse()
    if len(_PARSE_CACHE) < 20_000:
        _PARSE_CACHE[text] = node
    return node


# ---------------------------------------------------------------------------
#  Evaluation
# ---------------------------------------------------------------------------

def _num(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        if v.strip().upper() in _ERRORS:
            raise FormulaError(f"formula reaches an error value: {v}")
        try:
            return float(v.replace(",", "").strip())
        except ValueError as exc:
            raise FormulaError(f"cannot use {v!r} as a number") from exc
    raise FormulaError(f"cannot use {type(v).__name__} as a number")


def _text(v: Any) -> str:
    """Excel's text coercion: blank is "", whole numbers lose the .0, booleans go upper case."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _mid(text: str, start: int, count: int) -> str:
    if start < 1:
        raise FormulaError("MID: start position must be 1 or more")
    return text[start - 1: start - 1 + max(count, 0)]


def _flat(values: Iterable[Any]) -> List[Any]:
    out: List[Any] = []
    for v in values:
        out.extend(v) if isinstance(v, list) else out.append(v)
    return out


def _numbers(args: Sequence[Any]) -> List[float]:
    return [_num(v) for v in _flat(args)
            if v is not None and v != "" and not isinstance(v, str)]


def _div(a: float, b: float) -> float:
    if b == 0:
        raise FormulaError("division by zero")
    return a / b


FUNCTIONS: Dict[str, Callable[[List[Any]], Any]] = {
    "SUM":      lambda a: math.fsum(_numbers(a)),
    "AVERAGE":  lambda a: (math.fsum(_numbers(a)) / len(_numbers(a))) if _numbers(a) else 0.0,
    "MIN":      lambda a: min(_numbers(a)) if _numbers(a) else 0.0,
    "MAX":      lambda a: max(_numbers(a)) if _numbers(a) else 0.0,
    "COUNT":    lambda a: float(len(_numbers(a))),
    "ABS":      lambda a: abs(_num(a[0])),
    "SQRT":     lambda a: math.sqrt(_num(a[0])),
    "POWER":    lambda a: _num(a[0]) ** _num(a[1]),
    "ROUND":    lambda a: round(_num(a[0]), int(_num(a[1]))),
    "ROUNDUP":  lambda a: math.ceil(_num(a[0]) * 10 ** int(_num(a[1]))) / 10 ** int(_num(a[1])),
    "ROUNDDOWN": lambda a: math.floor(_num(a[0]) * 10 ** int(_num(a[1]))) / 10 ** int(_num(a[1])),
    "PRODUCT":  lambda a: math.prod(_numbers(a)) if _numbers(a) else 0.0,
    "SIGN":     lambda a: float((_num(a[0]) > 0) - (_num(a[0]) < 0)),
    # text - the vintage ROA formula pulls the month number out of the "M6" header
    "MID":      lambda a: _mid(_text(a[0]), int(_num(a[1])), int(_num(a[2]))),
    "LEFT":     lambda a: _text(a[0])[:int(_num(a[1])) if len(a) > 1 else 1],
    "RIGHT":    lambda a: _text(a[0])[-(int(_num(a[1])) if len(a) > 1 else 1):] or "",
    "LEN":      lambda a: float(len(_text(a[0]))),
    "TRIM":     lambda a: " ".join(_text(a[0]).split()),
    "UPPER":    lambda a: _text(a[0]).upper(),
    "LOWER":    lambda a: _text(a[0]).lower(),
    "VALUE":    lambda a: _num(_text(a[0])),
    "CONCATENATE": lambda a: "".join(_text(v) for v in _flat(a)),
    "CONCAT":   lambda a: "".join(_text(v) for v in _flat(a)),
    "N":        lambda a: _num(a[0]),
    "T":        lambda a: a[0] if isinstance(a[0], str) else "",
}
LAZY_FUNCTIONS = {"IF", "IFERROR", "IFNA"}          # evaluated by the walker, not the table above

_COMPARE = {
    "=":  lambda a, b: a == b,
    "<>": lambda a, b: a != b,
    "<":  lambda a, b: a < b,
    ">":  lambda a, b: a > b,
    "<=": lambda a, b: a <= b,
    ">=": lambda a, b: a >= b,
}


def _compare(op: str, a: Any, b: Any) -> bool:
    """Excel's rules, which are not Python's.

    An empty cell equals both "" and 0. Text and numbers are never equal to each other,
    and in an ordering comparison every number sorts below every piece of text.
    """
    if a is None and b is None:
        return _COMPARE[op](0, 0)
    if a is None:
        a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is None:
        b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)

    a_txt, b_txt = isinstance(a, str), isinstance(b, str)
    if a_txt and b_txt:
        return _COMPARE[op](a.casefold(), b.casefold())
    if a_txt != b_txt:                       # a number is never equal to text; numbers sort first
        if op == "=":
            return False
        if op == "<>":
            return True
        return _COMPARE[op](1 if a_txt else 0, 1 if b_txt else 0)
    return _COMPARE[op](_num(a), _num(b))


def evaluate(node: Node, resolve: Callable[[Optional[str], int, int], Any],
             sheet: Optional[str] = None) -> Any:
    """Evaluate an AST. `resolve(sheet, row, col)` returns one cell's value."""
    k = node.kind

    if k == "blank":
        return None
    if k in ("num", "str", "bool"):
        return node.value
    if k == "err":
        raise FormulaError(f"formula contains {node.value}")
    if k == "wholeref":
        raise UnsupportedFormula(
            f"whole-column/row reference {node.value!r} is not supported - "
            "the trace needs a bounded range like B10:B248")
    if k == "ref":
        r, c = node.coord
        return resolve(node.sheet or sheet, r, c)
    if k == "range":
        (r1, c1), (r2, c2) = node.coord
        sh = node.sheet or sheet
        return [resolve(sh, r, c)
                for r in range(min(r1, r2), max(r1, r2) + 1)
                for c in range(min(c1, c2), max(c1, c2) + 1)]
    if k == "list":
        return [evaluate(a, resolve, sheet) for a in node.args]

    if k == "unary":
        if node.value == "%":
            return _num(evaluate(node.args[0], resolve, sheet)) / 100.0
        v = _num(evaluate(node.args[0], resolve, sheet))
        return -v if node.value == "-" else v

    if k == "binop":
        op = node.value
        if op == "&":
            return _text(evaluate(node.args[0], resolve, sheet)) + \
                   _text(evaluate(node.args[1], resolve, sheet))
        a = evaluate(node.args[0], resolve, sheet)
        b = evaluate(node.args[1], resolve, sheet)
        if op in _COMPARE:
            return _compare(op, a, b)
        x, y = _num(a), _num(b)
        if op == "+": return x + y
        if op == "-": return x - y
        if op == "*": return x * y
        if op == "/": return _div(x, y)
        if op == "^": return x ** y
        raise UnsupportedFormula(f"operator {op!r}")

    if k == "func":
        name = node.value
        if name == "IF":
            cond = evaluate(node.args[0], resolve, sheet)
            truthy = bool(cond) if isinstance(cond, bool) else _num(cond) != 0
            if truthy:
                return evaluate(node.args[1], resolve, sheet)
            return evaluate(node.args[2], resolve, sheet) if len(node.args) > 2 else False
        if name in ("IFERROR", "IFNA"):
            try:
                return evaluate(node.args[0], resolve, sheet)
            except UnsupportedFormula:
                raise                       # never let IFERROR hide a gap in this evaluator
            except FormulaError:
                return evaluate(node.args[1], resolve, sheet)
        fn = FUNCTIONS.get(name)
        if fn is None:
            raise UnsupportedFormula(
                f"{name}() is not implemented. Supported: "
                f"{', '.join(sorted(set(FUNCTIONS) | LAZY_FUNCTIONS))}.")
        return fn([evaluate(a, resolve, sheet) for a in node.args])

    raise UnsupportedFormula(f"node kind {k!r}")


def refs_in(node: Node, sheet: Optional[str] = None) -> List[Tuple[Optional[str], int, int]]:
    """Every single cell the formula touches, ranges expanded."""
    out: List[Tuple[Optional[str], int, int]] = []
    if node.kind == "ref":
        out.append((node.sheet or sheet, *node.coord))
    elif node.kind == "range":
        (r1, c1), (r2, c2) = node.coord
        sh = node.sheet or sheet
        out += [(sh, r, c)
                for r in range(min(r1, r2), max(r1, r2) + 1)
                for c in range(min(c1, c2), max(c1, c2) + 1)]
    for a in node.args:
        out += refs_in(a, sheet)
    return out


# ---------------------------------------------------------------------------
#  Workbooks: cached values and formula text, side by side
# ---------------------------------------------------------------------------

@dataclass
class Book:
    path: Path
    values: Dict[str, Dict[Tuple[int, int], Any]]
    formulas: Dict[str, Dict[Tuple[int, int], str]]
    max_row: Optional[int] = None        # the window that was read, so we can spot refs outside it
    max_col: Optional[int] = None

    def value(self, sheet: str, row: int, col: int) -> Any:
        return self.values.get(sheet, {}).get((row, col))

    def formula(self, sheet: str, row: int, col: int) -> Optional[str]:
        return self.formulas.get(sheet, {}).get((row, col))

    @property
    def sheets(self) -> List[str]:
        return list(self.values)


def sheet_names(path: str | Path) -> List[str]:
    """Just the sheet names - cheap, so you can decide what is worth reading."""
    wb = load_workbook(Path(path), read_only=True, data_only=True, keep_links=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_book(path: str | Path, sheets: Optional[Iterable[str]] = None,
              max_row: Optional[int] = None, max_col: Optional[int] = None) -> Book:
    """Read cached values and formula text from an .xlsx/.xlsm.

    `sheets`, `max_row` and `max_col` bound the read. A 46-sheet report where you only need
    24 sheets and the first 260 rows is a fraction of the work - and a reference landing
    outside the window raises rather than quietly reading as blank.
    """
    p = Path(path)
    if p.suffix.lower() in {".xlsb", ".xls"}:
        raise UnsupportedFormula(
            f"{p.name}: formulas cannot be read from '{p.suffix}' - pyxlsb and xlrd only expose "
            "cached values. Save the workbook as .xlsx (File > Save As > Excel Workbook) to trace it. "
            "The value comparison works on the original file either way.")
    wanted = set(sheets) if sheets is not None else None
    values: Dict[str, Dict[Tuple[int, int], Any]] = {}
    formulas: Dict[str, Dict[Tuple[int, int], str]] = {}

    for data_only, sink in ((True, values), (False, formulas)):
        wb = load_workbook(p, data_only=data_only, read_only=True, keep_links=False)
        try:
            for name in wb.sheetnames:
                if wanted is not None and name not in wanted:
                    continue
                ws = wb[name]
                if not hasattr(ws, "iter_rows"):
                    continue
                grid: Dict[Tuple[int, int], Any] = {}
                for row in ws.iter_rows(max_row=max_row, max_col=max_col):
                    for cell in row:
                        v = cell.value
                        if v is None:
                            continue
                        if data_only:
                            grid[(cell.row, cell.column)] = v
                        elif isinstance(v, str) and v.startswith("="):
                            grid[(cell.row, cell.column)] = v
                        elif type(v).__name__ in ("ArrayFormula", "DataTableFormula"):
                            grid[(cell.row, cell.column)] = getattr(v, "text", "=?")
                sink[name] = grid
        finally:
            wb.close()
    return Book(path=p.resolve(), values=values, formulas=formulas,
                max_row=max_row, max_col=max_col)


# ---------------------------------------------------------------------------
#  Tracing
# ---------------------------------------------------------------------------

@dataclass
class Component:
    key: str                    # stable id, e.g. "Vintage 1!C168"
    sheet: str
    row_old: int
    row_new: int
    col: int
    label: str                  # from the label column, or the cell ref when there isn't one
    old: Optional[float]
    new: Optional[float]
    depth: int
    labelled: bool
    formula: Optional[str] = None

    @property
    def ref_old(self) -> str:
        return f"{get_column_letter(self.col)}{self.row_old}"

    @property
    def ref_new(self) -> str:
        return f"{get_column_letter(self.col)}{self.row_new}"

    @property
    def delta(self) -> float:
        return (self.new or 0.0) - (self.old or 0.0)

    @property
    def blank_side(self) -> str:
        if self.old is None:
            return "old"
        return "new" if self.new is None else ""


@dataclass
class StructuralChange:
    """A cell whose formula differs between the workbooks once the row shift is allowed for."""
    sheet: str
    ref_old: str
    ref_new: str
    formula_old: str
    formula_new: str
    label: str = ""
    kind: str = "formula rewritten"
    extra_rows: List[int] = field(default_factory=list)   # rows the new range covers and the old did not
    extra_detail: str = ""

    def describe(self) -> str:
        if self.kind == "range widened":
            return (f"{self.label or self.ref_old}: the range now also covers "
                    f"row{'s' if len(self.extra_rows) > 1 else ''} "
                    f"{', '.join(str(r) for r in self.extra_rows)} of the new workbook"
                    + (f" ({self.extra_detail})" if self.extra_detail else ""))
        return f"{self.label or self.ref_old}: formula rewritten"


@dataclass
class Trace:
    sheet: str
    col: int
    row_old: int
    row_new: int
    formula_old: Optional[str]
    formula_new: Optional[str]
    value_old: float
    value_new: float
    components: List[Component]
    tree: Node                  # the formula, rebuilt with components as parameters
    notes: List[str] = field(default_factory=list)
    structural: List[StructuralChange] = field(default_factory=list)
    modelled_new: Optional[float] = None   # old formula driven by all-new inputs

    @property
    def ref_old(self) -> str:
        return f"{get_column_letter(self.col)}{self.row_old}"

    @property
    def ref_new(self) -> str:
        return f"{get_column_letter(self.col)}{self.row_new}"

    @property
    def formula_changed(self) -> bool:
        return (self.formula_old or "").replace(" ", "") != (self.formula_new or "").replace(" ", "")

    _folded: Optional[Node] = field(default=None, repr=False, compare=False)

    @property
    def folded(self) -> Node:
        """The tree with every component-free subtree pre-computed. Built once."""
        if self._folded is None:
            self._folded = fold_constants(self.tree)
        return self._folded

    def evaluate_with(self, selected: Set[str]) -> float:
        """Re-evaluate with the named components switched to their new values."""
        picked = {c.key: (c.new if c.key in selected else c.old) for c in self.components}
        return float(eval_dag(self.folded, picked))

    @property
    def structural_gap(self) -> Optional[float]:
        """New ROA the workbook reports, less what the old formula makes of all-new inputs.

        Anything other than zero means the calculation itself changed shape - a formula
        rewritten, or a range that now spans an inserted row - rather than only its inputs
        moving. That part of the move cannot be attributed to any component.
        """
        if self.value_new is None or self.modelled_new is None:
            return None
        return self.value_new - self.modelled_new

    def check(self, tol: float = 1e-6) -> None:
        """The old side must reconcile exactly, or the formula was followed incorrectly.

        The new side is *reported*, not asserted: a gap there is a real finding about the
        workbooks (the calculation changed shape), not a bug in the trace.
        """
        if self.value_old is not None:
            got = self.evaluate_with(set())
            if abs(got - self.value_old) / max(abs(self.value_old), 1.0) > tol:
                raise AssertionError(
                    f"trace does not reproduce the old value at {self.sheet}!{self.ref_old}: "
                    f"re-evaluated {got!r}, workbook says {self.value_old!r}. The formula was "
                    "followed incorrectly - do not trust the component impacts.")


_MISS = object()


def eval_dag(node: Node, picked: Dict[str, Any], memo: Optional[Dict[int, Any]] = None) -> Any:
    """Evaluate the scenario graph, visiting each distinct node once.

    The rebuilt formula is a DAG, not a tree: a cell mentioned twice in one formula - your
    IFERROR(IF(D178-D202=0,"",D178-D202),"") mentions D202 twice - is one node with two
    parents. Walking it as a tree re-expands it, and the cost doubles at every level: a
    50-node graph became hundreds of millions of visits. Memoising on node identity makes
    it linear, which is the difference between hours and milliseconds.
    """
    if memo is None:
        memo = {}
    key = id(node)
    hit = memo.get(key, _MISS)
    if hit is not _MISS:
        return hit

    k = node.kind
    if k == "param":
        v = picked[node.value]
    elif k in ("num", "str", "bool"):
        v = node.value
    elif k == "blank":
        v = None
    elif k == "err":
        raise FormulaError(f"formula contains {node.value}")
    elif k == "list":
        v = [eval_dag(a, picked, memo) for a in node.args]
    elif k == "unary":
        if node.value == "%":
            v = _num(eval_dag(node.args[0], picked, memo)) / 100.0
        else:
            x = _num(eval_dag(node.args[0], picked, memo))
            v = -x if node.value == "-" else x
    elif k == "binop":
        op = node.value
        if op == "&":
            v = _text(eval_dag(node.args[0], picked, memo)) + \
                _text(eval_dag(node.args[1], picked, memo))
        else:
            a = eval_dag(node.args[0], picked, memo)
            b = eval_dag(node.args[1], picked, memo)
            if op in _COMPARE:
                v = _compare(op, a, b)
            else:
                x, y = _num(a), _num(b)
                if op == "+":   v = x + y
                elif op == "-": v = x - y
                elif op == "*": v = x * y
                elif op == "/": v = _div(x, y)
                elif op == "^": v = x ** y
                else: raise UnsupportedFormula(f"operator {op!r}")
    elif k == "func":
        name = node.value
        if name == "IF":
            cond = eval_dag(node.args[0], picked, memo)
            truthy = bool(cond) if isinstance(cond, bool) else (cond is not None and _num(cond) != 0)
            if truthy:
                v = eval_dag(node.args[1], picked, memo)
            else:
                v = eval_dag(node.args[2], picked, memo) if len(node.args) > 2 else False
        elif name in ("IFERROR", "IFNA"):
            try:
                v = eval_dag(node.args[0], picked, memo)
            except UnsupportedFormula:
                raise
            except FormulaError:
                v = eval_dag(node.args[1], picked, memo)
        else:
            fn = FUNCTIONS.get(name)
            if fn is None:
                raise UnsupportedFormula(f"{name}() is not implemented")
            v = fn([eval_dag(a, picked, memo) for a in node.args])
    else:
        raise UnsupportedFormula(f"node kind {k!r}")

    memo[key] = v
    return v


def _eval_scenario(node: Node, picked: Dict[str, Any]) -> Any:
    return eval_dag(node, picked)


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _blankish(v: Any) -> bool:
    """Excel treats an empty cell and a formula returning "" the same way. So do we."""
    return v is None or (isinstance(v, str) and v.strip() == "")


def remap_rows(text: str, rmap: Callable[[int], int]) -> str:
    """Rewrite the row numbers in a formula's references through the row map.

    Lets an old formula be compared against the new one on equal terms: if the only
    difference is the inserted row, the remapped old text and the new text match.
    String literals are left alone, so "M1" is never mistaken for a reference.
    """
    src = text[1:] if text.startswith("=") else text
    out, i = [], 0
    for tok in tokenize(src):
        out.append(src[i:tok.pos])
        if tok.kind in ("ref", "range"):
            def sub(m):
                return f"{m.group(1)}{rmap(int(m.group(2)))}"
            out.append(re.sub(r"(\$?[A-Za-z]{1,3}\$?)(\d{1,7})", sub, tok.text))
        else:
            out.append(tok.text)
        i = tok.pos + len(tok.text)
    out.append(src[i:])
    return "=" + "".join(out)


def _same(a: Any, b: Any, tol: float) -> bool:
    if _blankish(a) and _blankish(b):
        return True
    if _is_num(a) and _is_num(b):
        return abs(float(a) - float(b)) <= tol
    if _blankish(a) != _blankish(b):
        return False
    return str(a).strip().casefold() == str(b).strip().casefold()


def trace(old: Book, new: Book, sheet: str, row: int, col: int,
          row_map: Optional[Callable[[int], int]] = None,
          label_col: int = 2, max_depth: int = 8, tol: float = 1e-9,
          sheet_new: Optional[str] = None, max_cells: int = 20_000) -> Trace:
    """Follow the formula at (sheet, row, col) down to the inputs that changed.

    `row_map` maps a row in the old workbook to the same logical row in the new one - this is
    what carries the inserted-row offset, so old row 209 can be compared against new row 210.

    A cell becomes a *component* when it changed and nothing labelled below it changed: that is
    the individual revenue or expense line that moved, however many subtotal layers sit above it.
    """
    rmap = row_map or (lambda r: r)
    sheet_b = sheet_new or sheet
    notes: List[str] = []
    components: Dict[str, Component] = {}
    structural: List[StructuralChange] = []
    memo: Dict[Tuple[str, int, int], Tuple[Node, bool]] = {}
    visiting: Set[Tuple[str, int, int]] = set()
    entered = [0]          # cells ENTERED, not cells finished: the memo only fills on the
                           # way back up, so counting finished cells never stops a deep descent

    def label_of(sh: str, r: int) -> str:
        v = old.value(sh, r, label_col)
        return str(v).strip() if v is not None and str(v).strip() else ""

    def add_component(sh: str, r: int, c: int, ov: Any, nv: Any, depth: int, lab: str) -> Node:
        key = f"{sh}!{get_column_letter(c)}{r}"
        if key not in components:
            components[key] = Component(
                key=key, sheet=sh, row_old=r, row_new=rmap(r), col=c,
                label=lab or f"{get_column_letter(c)}{r}",
                old=None if ov is None else float(ov),
                new=None if nv is None else float(nv),
                depth=depth, labelled=bool(lab), formula=old.formula(sh, r, c))
        return Node("param", key)

    def rebuild(node: Node, sh: str, depth: int) -> Tuple[Node, bool]:
        if node.kind == "ref":
            r, c = node.coord
            return visit(node.sheet or sh, r, c, depth + 1)
        if node.kind == "range":
            (r1, c1), (r2, c2) = node.coord
            target = node.sheet or sh
            kids, found = [], False
            for r in range(min(r1, r2), max(r1, r2) + 1):
                for c in range(min(c1, c2), max(c1, c2) + 1):
                    kid, f = visit(target, r, c, depth + 1)
                    kids.append(kid)
                    found = found or f
            return Node("list", None, kids), found
        kids, found = [], False
        for a in node.args:
            kid, f = rebuild(a, sh, depth)
            kids.append(kid)
            found = found or f
        return Node(node.kind, node.value, kids, node.sheet, node.coord), found

    def norm_formula(t: str) -> str:
        return re.sub(r"\s+", "", t or "").upper()

    def ranges_of(node: Node) -> List[Node]:
        out = [node] if node.kind == "range" else []
        for a in node.args:
            out += ranges_of(a)
        return out

    def note_structural(sh: str, r: int, c: int, sh_new: str, f_old, f_new) -> None:
        if f_old is None and f_new is None:
            return
        ref_old = f"{get_column_letter(c)}{r}"
        ref_new = f"{get_column_letter(c)}{rmap(r)}"
        lab = label_of(sh, r)

        if f_old is None or f_new is None:
            structural.append(StructuralChange(
                sheet=sh, ref_old=ref_old, ref_new=ref_new,
                formula_old=f_old or "(no formula - a plain value)",
                formula_new=f_new or "(no formula - a plain value)", label=lab))
            return

        try:
            remapped = remap_rows(f_old, rmap)
        except FormulaError:
            remapped = f_old
        if norm_formula(remapped) != norm_formula(f_new):
            structural.append(StructuralChange(
                sheet=sh, ref_old=ref_old, ref_new=ref_new,
                formula_old=f_old, formula_new=f_new, label=lab))
            return

        # The texts match after remapping - but a range that SPANS the inserted row grows by a
        # row while reading identically, so it silently picks up a cell with no old counterpart.
        try:
            r_old, r_new = ranges_of(parse(f_old)), ranges_of(parse(f_new))
        except FormulaError:
            return
        if len(r_old) != len(r_new):
            return
        for go, gn in zip(r_old, r_new):
            (o1, oc1), (o2, oc2) = go.coord
            (n1, nc1), (n2, nc2) = gn.coord
            if {oc1, oc2} != {nc1, nc2}:
                continue
            covered = {rmap(x) for x in range(min(o1, o2), max(o1, o2) + 1)}
            extra = sorted(set(range(min(n1, n2), max(n1, n2) + 1)) - covered)
            if not extra:
                continue
            bits = []
            for x in extra[:4]:
                name = new.value(sh_new, x, label_col)
                val = new.value(sh_new, x, min(nc1, nc2))
                bits.append(f"row {x} {str(name).strip()!r}" + (f" = {val:,.2f}" if _is_num(val) else ""))
            structural.append(StructuralChange(
                sheet=sh, ref_old=ref_old, ref_new=ref_new,
                formula_old=f_old, formula_new=f_new, label=lab,
                kind="range widened", extra_rows=extra, extra_detail="; ".join(bits)))

    def visit(sh: str, r: int, c: int, depth: int) -> Tuple[Node, bool]:
        key = (sh, r, c)
        if key in memo:
            return memo[key]
        if key in visiting:
            raise FormulaError(f"circular reference at {sh}!{get_column_letter(c)}{r}")

        entered[0] += 1
        sh_new = sheet_b if sh == sheet else sh
        if (old.max_row and r > old.max_row) or (old.max_col and c > old.max_col):
            raise FormulaError(
                f"{sh}!{get_column_letter(c)}{r} is outside the window the workbook was read with "
                f"(max_row={old.max_row}, max_col={old.max_col}). Widen it in load_book, or the "
                "cell would read as blank and the answer would be wrong.")
        ov = old.value(sh, r, c)
        nv = new.value(sh_new, rmap(r), c)
        lab = label_of(sh, r)
        f = old.formula(sh, r, c)
        note_structural(sh, r, c, sh_new, f, new.formula(sh_new, rmap(r), c))

        if _same(ov, nv, tol):
            # Keep the value's TYPE, not just its number. Formulas branch on text -
            # IF(D5="M1", ...) reads the month header - so a text constant that happens
            # not to have changed must still be carried through as text.
            if _is_num(ov):
                result = (Node("num", float(ov)), False)
            elif isinstance(ov, bool):
                result = (Node("bool", ov), False)
            elif isinstance(ov, str) and not _blankish(ov):
                result = (Node("str", ov), False)
            elif f is not None and depth < max_depth and entered[0] < max_cells:
                # A workbook with no cached values reaches every cell through here, so the
                # ceiling has to guard this path too, not only the changed one.
                visiting.add(key)
                try:
                    sub, _ = rebuild(parse(f), sh, depth)
                    result = (sub, False)
                except FormulaError as exc:
                    notes.append(f"{sh}!{get_column_letter(c)}{r}: {exc} - held blank")
                    result = (Node("blank"), False)
                finally:
                    visiting.discard(key)
            elif _blankish(ov):
                if f is not None and entered[0] >= max_cells:
                    notes.append(f"stopped at {sh}!{get_column_letter(c)}{r} "
                                 f"(cell limit {max_cells:,}) - held blank")
                result = (Node("blank"), False)          # blank, not zero: "" comparisons depend on it
            else:
                notes.append(f"{sh}!{get_column_letter(c)}{r} holds {ov!r}, not a number - held blank")
                result = (Node("blank"), False)
            memo[key] = result
            return result

        ov_n = None if _blankish(ov) else ov
        nv_n = None if _blankish(nv) else nv
        if not (_is_num(ov_n) or ov_n is None) or not (_is_num(nv_n) or nv_n is None):
            notes.append(f"{sh}!{get_column_letter(c)}{r} changed from {ov!r} to {nv!r} - "
                         "not a number, so it is held at its old value in the model")
            memo[key] = (Node("blank") if ov_n is None else Node("num", float(ov_n)), False)
            return memo[key]

        if f is None or depth >= max_depth or entered[0] >= max_cells:
            if f is not None:
                why = (f"depth limit {max_depth}" if depth >= max_depth
                       else f"cell limit {max_cells:,}")
                notes.append(f"stopped at {sh}!{get_column_letter(c)}{r} ({why}) - "
                             "anything below it is folded into this component")
            result = (add_component(sh, r, c, ov_n, nv_n, depth, lab), bool(lab))
            memo[key] = result
            return result

        visiting.add(key)
        try:
            sub, found_below = rebuild(parse(f), sh, depth)
        except FormulaError as exc:
            notes.append(f"{sh}!{get_column_letter(c)}{r}: {exc} - treated as a single component")
            visiting.discard(key)
            result = (add_component(sh, r, c, ov_n, nv_n, depth, lab), bool(lab))
            memo[key] = result
            return result
        visiting.discard(key)

        if lab and not found_below:
            result = (add_component(sh, r, c, ov_n, nv_n, depth, lab), True)
            memo[key] = result
            return result

        result = (sub, found_below or bool(lab))
        memo[key] = result
        return result

    f_old = old.formula(sheet, row, col)
    f_new = new.formula(sheet_b, rmap(row), col)
    if f_old is None:
        raise FormulaError(
            f"{sheet}!{get_column_letter(col)}{row} holds no formula - it is a plain value, so "
            "there is nothing to decompose. Check the row and column.")

    tree, _ = rebuild(parse(f_old), sheet, 0)

    v_old = old.value(sheet, row, col)
    v_new = new.value(sheet_b, rmap(row), col)
    result = Trace(
        sheet=sheet, col=col, row_old=row, row_new=rmap(row),
        formula_old=f_old, formula_new=f_new,
        value_old=float(v_old) if _is_num(v_old) else None,
        value_new=float(v_new) if _is_num(v_new) else None,
        components=sorted(components.values(), key=lambda c: -abs(c.delta)),
        tree=tree, notes=notes, structural=structural)
    result.modelled_new = result.evaluate_with({c.key for c in result.components})

    if result.value_old is None or result.value_new is None:
        notes.append(
            f"{sheet} carries no cached value at the ROA cell, so both figures below are "
            "computed from the formula rather than read from the workbook. Open the file in "
            "Excel and re-save it to store calculated values and cross-check them.")
        if result.value_old is None:
            result.value_old = result.evaluate_with(set())
        if result.value_new is None:
            result.value_new = result.evaluate_with({c.key for c in result.components})
    widened = [sc for sc in structural if sc.kind == "range widened"]
    if widened:
        result.notes.insert(0, (
            f"{len(widened)} range(s) in the new workbook cover rows the old one does not - the "
            "inserted row falls inside them, so those cells feed the new figure and have no "
            "counterpart to be compared against. " + " ".join(sc.describe() for sc in widened[:3])))

    gap = result.structural_gap
    if gap is not None and abs(gap) > 1e-9:
        result.notes.insert(0, (
            f"{abs(gap) * 10_000:,.1f} bps of the move does not come from any component. The "
            f"calculation itself changed shape between the workbooks - see the {len(structural)} "
            "formula difference(s) listed. Driving the old formula with every new input reaches "
            f"{result.modelled_new:.6%}, but the snow workbook reports {result.value_new:.6%}."))
    return result


def _has_param(n: Node) -> bool:
    return n.kind == "param" or any(_has_param(a) for a in n.args)


def fold_constants(node: Node, _cache: Optional[Dict[int, Node]] = None) -> Node:
    """Collapse every subtree that contains no component into a single literal.

    Nearly all of a real formula tree is machinery that did not change - the months that
    were already equal, the untouched expense lines. Pre-computing those turns thousands
    of nodes into one number each, which is what keeps the page a sensible size.

    The cache keeps one input node mapping to one output node, so the sharing the walk
    established survives folding and the serialiser can still emit repeats by reference.
    """
    if _cache is None:
        _cache = {}
    hit = _cache.get(id(node))
    if hit is not None:
        return hit
    if node.kind in ("param", "num", "str", "bool", "blank", "err"):
        _cache[id(node)] = node
        return node
    folded = Node(node.kind, node.value, [fold_constants(a, _cache) for a in node.args],
                  node.sheet, node.coord)
    _cache[id(node)] = folded
    if node.kind not in ("func", "binop", "unary") or _has_param(folded):
        return folded
    try:
        v = eval_dag(folded, {})
    except FormulaError:
        return folded                      # an error an enclosing IFERROR is meant to catch
    if v is None:
        lit = Node("blank")
    elif isinstance(v, bool):
        lit = Node("bool", v)
    elif isinstance(v, (int, float)):
        lit = Node("num", float(v))
    elif isinstance(v, str):
        lit = Node("str", v)
    else:
        return folded
    _cache[id(node)] = lit
    return lit


def compile_scenario(tr: Trace) -> Dict[str, Any]:
    """Serialise the scenario tree so a browser can re-evaluate it.

    The walk memoises, so one cell reached by several paths is one Node object. Those are
    emitted once into `defs` and referenced, instead of being written out at every
    occurrence - a formula like IF(a-b=0,"",a-b) mentions the same subtree twice.
    """
    defs: List[Dict[str, Any]] = []
    index: Dict[int, int] = {}
    seen: Dict[int, int] = {}

    def pack(n: Node) -> Dict[str, Any]:
        if n.kind == "param":
            return {"k": "p", "id": n.value}
        if n.kind == "blank":
            return {"k": "z"}
        if n.kind in ("num", "bool"):
            return {"k": "n", "v": n.value}
        if n.kind == "str":
            return {"k": "s", "v": n.value}
        if n.kind == "list":
            return {"k": "l", "a": [pack(a) for a in n.args]}
        if n.kind == "binop":
            return {"k": "b", "o": n.value, "a": [pack(a) for a in n.args]}
        if n.kind == "unary":
            return {"k": "u", "o": n.value, "a": [pack(a) for a in n.args]}
        if n.kind == "func":
            return {"k": "f", "o": n.value, "a": [pack(a) for a in n.args]}
        raise UnsupportedFormula(f"cannot serialise node {n.kind!r}")

    def share(n: Node) -> Dict[str, Any]:
        if n.kind in ("num", "str", "bool", "blank", "err", "param"):
            return pack(n)
        key = id(n)
        if key in index:
            return {"k": "r", "i": index[key]}
        if seen.get(key):
            slot = len(defs)
            defs.append(None)                       # reserve, then fill after packing children
            index[key] = slot
            defs[slot] = pack_shared(n)
            return {"k": "r", "i": slot}
        seen[key] = 1
        return pack_shared(n)

    def pack_shared(n: Node) -> Dict[str, Any]:
        if n.kind == "list":
            return {"k": "l", "a": [share(a) for a in n.args]}
        if n.kind == "binop":
            return {"k": "b", "o": n.value, "a": [share(a) for a in n.args]}
        if n.kind == "unary":
            return {"k": "u", "o": n.value, "a": [share(a) for a in n.args]}
        if n.kind == "func":
            return {"k": "f", "o": n.value, "a": [share(a) for a in n.args]}
        return pack(n)

    def count(n: Node) -> None:
        seen[id(n)] = seen.get(id(n), 0) + 1
        if seen[id(n)] == 1:
            for a in n.args:
                count(a)

    folded = tr.folded
    count(folded)
    repeated = {k for k, v in seen.items() if v > 1}
    seen = {k: 1 for k in repeated}                 # only share what actually repeats
    root = share(folded)
    return {"defs": defs, "root": root}


def trace_columns(old: Book, new: Book, sheet: str, row: int, cols: Iterable[int],
                  **kw) -> Tuple[List["Trace"], List[Tuple[str, str]]]:
    """Trace the same row across many columns - one per month, in the vintage sheets.

    Returns (traces, failures). A column that cannot be traced is reported with its reason
    rather than dropping out of the result silently.
    """
    traces: List[Trace] = []
    failures: List[Tuple[str, str]] = []
    for c in cols:
        ref = f"{get_column_letter(c)}{row}"
        if old.formula(sheet, row, c) is None:
            failures.append((ref, "no formula in the old workbook"))
            continue
        try:
            tr = trace(old, new, sheet, row, c, **kw)
            tr.check()
        except (FormulaError, AssertionError) as exc:
            failures.append((ref, str(exc)))
            continue
        traces.append(tr)
    return traces, failures
