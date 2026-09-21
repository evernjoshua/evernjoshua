"""
Finance Impact Notes - month-over-month diff of the P&L forecast tabs, formatted
to match an existing "Loss Notes" sheet.

Script export of finance_impact.ipynb; the notebook is the primary copy and this
is generated from its code cells, so edit the notebook and re-export rather than
editing here. See the notebook's closing notes for the scale/percentage caveat and
for what the template harvester does and does not carry over.
"""

import os
import re
from copy import copy
from dataclasses import dataclass, field
from statistics import median
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------- Source workbooks ----------
CURR_WORKBOOK = r"\\fnbm.corp\share\Risk\Portfolio Growth\Personal Folders\Lucas Lin\Digital\PAF\2026\10'26\10'26 P&L Forecast - All.xlsm"
PREV_WORKBOOK = r"\\fnbm.corp\share\Risk\Portfolio Growth\Personal Folders\Lucas Lin\Digital\PAF\2026\09'26\09'26 P&L Forecast - All.xlsm"
OUTPUT_PATH   = r"\\fnbm.corp\share\Risk\Portfolio Growth\Personal Folders\Lucas Lin\Digital\PAF\2026\10'26\ROA Decomp & Finance Impact Notes\finance_impact.xlsx"

# ---------- Formatting template ----------
# The sheet whose look-and-feel gets cloned onto the output.
TEMPLATE_WORKBOOK = r"\\fnbm.corp\share\Risk\Portfolio Growth\Personal Folders\Lucas Lin\Multiple Account Program\PAF\2026\06'26\$800 CL Scenario\Financial Impact\06'26 Finance Impact Notes.xlsx"
TEMPLATE_SHEET    = "Loss Notes"

# ---------- Which product tabs to walk ----------
START_SHEET   = "MoneyLion $75 601-6"
END_SHEET     = "Organic $95 <=600"
SUMMARY_SHEET = "summary"

# ---------- Metrics ----------
# row           -> worksheet row the metric sits on in each product tab
# scale         -> multiplier on the delta. Leave at 1 to keep raw cell units.
#                  Set to 100 if the cells hold true percentages (0.0525) and you
#                  want percentage-point deltas (5.25) in the output.
# number_format -> Excel format for that metric's three delta columns
DELTA_FORMAT = "+0.00;-0.00;0.00"   # + when up, - when down, 0.00 when flat

METRICS = [
    {"label": "Unit C/O",            "row": 11, "scale": 1, "number_format": DELTA_FORMAT},
    {"label": "$ C/O",               "row": 12, "scale": 1, "number_format": DELTA_FORMAT},
    {"label": "Gross Revenue",       "row": 15, "scale": 1, "number_format": DELTA_FORMAT},
    {"label": "Provision",           "row": 16, "scale": 1, "number_format": DELTA_FORMAT},
    {"label": "Total Expenses",      "row": 18, "scale": 1, "number_format": DELTA_FORMAT},
    {"label": "Cuml ROA Annualized", "row": 22, "scale": 1, "number_format": DELTA_FORMAT},
]

YEAR_COLUMNS = ["H", "I", "J"]                   # Yr 1, Yr 2, Yr 3 on the source tabs
YEAR_LABELS  = ["Yr 1", "2Yr Cuml", "3Yr Cuml"]  # sub-header under each metric

ID_HEADERS   = ["Product", "First Word", "Vantage", "Credit Line"]
NOTE_HEADERS = ["Credit Line Notes", "Product Notes"]

# ---------- Formatting switches ----------
BAND_BY                   = "group"   # "group" = one shade per Product/Vantage block, "row" = every other row
USE_TEMPLATE_NUMBER_FORMAT = False    # True = inherit Loss Notes' number format instead of DELTA_FORMAT
COPY_COLUMN_WIDTHS         = True
COPY_ROW_HEIGHTS           = True
COPY_TAB_COLOR             = True
COPY_GRIDLINE_SETTING      = True

FIRST_DELTA_COL = len(ID_HEADERS) + 1
DELTA_COL_COUNT = len(METRICS) * len(YEAR_COLUMNS)
LAST_DELTA_COL  = FIRST_DELTA_COL + DELTA_COL_COUNT - 1
FIRST_NOTE_COL  = LAST_DELTA_COL + 1
TOTAL_COLS      = LAST_DELTA_COL + len(NOTE_HEADERS)

HEADER_ROWS    = 2                  # row 1 = metric name, row 2 = year
DATA_START_ROW = HEADER_ROWS + 1

# column index -> the metric it belongs to
METRIC_BY_COL: Dict[int, dict] = {}
_col = FIRST_DELTA_COL
for _metric in METRICS:
    for _offset in range(len(YEAR_COLUMNS)):
        METRIC_BY_COL[_col + _offset] = _metric
    _col += len(YEAR_COLUMNS)

# columns that open a new metric block (heavier left edge)
BOUNDARY_COLS = set(range(FIRST_DELTA_COL, LAST_DELTA_COL + 1, len(YEAR_COLUMNS))) | {FIRST_NOTE_COL}

print(f"{len(METRICS)} metrics x {len(YEAR_COLUMNS)} years = {DELTA_COL_COUNT} delta columns")
print(f"deltas in {get_column_letter(FIRST_DELTA_COL)}:{get_column_letter(LAST_DELTA_COL)}, "
      f"notes in {get_column_letter(FIRST_NOTE_COL)}:{get_column_letter(TOTAL_COLS)}, "
      f"data starts row {DATA_START_ROW}")

def parse_sheet_name(name: str):
    """
    'MoneyLion $75 601-6'          -> ('MoneyLion', '$75', '601-6', '')
    'Credit Karma $95_$400 <=600'  -> ('Credit Karma', '$95', '<=600', '$400')

    first_word  = everything before the first '$'
    product     = first dollar amount
    credit_line = second dollar amount, if there is one
    vantage     = last whitespace-delimited token
    """
    name = name.strip()
    first_word = name.split("$")[0].strip() if "$" in name else name

    dollars = re.findall(r"\$\d[\d,]*(?:\.\d+)?", name)
    product     = dollars[0] if len(dollars) >= 1 else ""
    credit_line = dollars[1] if len(dollars) >= 2 else ""
    vantage     = name.rsplit(" ", 1)[-1] if " " in name else ""
    return first_word, product, vantage, credit_line


def to_number(value):
    """Best-effort float. Handles $, commas, %, and (123) negatives. None if unparseable."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).replace("$", "").replace(",", "").strip()
    if s.endswith("%"):
        s = s[:-1].strip()

    negative = s.startswith("(") and s.endswith(")")
    if negative:
        s = s[1:-1]

    try:
        return -float(s) if negative else float(s)
    except ValueError:
        return None


def compare_values(curr_val, prev_val, scale=1):
    """(curr - prev) * scale, rounded. None if either side is missing."""
    a, b = to_number(curr_val), to_number(prev_val)
    if a is None or b is None:
        return None
    return round((a - b) * scale, 4)


def collect_deltas(curr_sheet, prev_sheet) -> List[Optional[float]]:
    """Flat list of deltas for one product tab: metric by metric, year by year."""
    deltas = []
    for metric in METRICS:
        for col in YEAR_COLUMNS:
            if prev_sheet is None:
                deltas.append(None)
                continue
            ref = f"{col}{metric['row']}"
            deltas.append(compare_values(curr_sheet[ref].value,
                                         prev_sheet[ref].value,
                                         metric.get("scale", 1)))
    return deltas

@dataclass
class StyleSnapshot:
    """One cell's formatting, detached from its source workbook."""
    font: Font
    fill: PatternFill
    border: Border
    alignment: Alignment
    number_format: str = "General"

    @classmethod
    def from_cell(cls, cell) -> "StyleSnapshot":
        return cls(font=copy(cell.font),
                   fill=copy(cell.fill),
                   border=copy(cell.border),
                   alignment=copy(cell.alignment),
                   number_format=cell.number_format)

    def apply(self, cell, number_format: Optional[str] = None,
              fill: Optional[PatternFill] = None,
              alignment: Optional[Alignment] = None,
              border: Optional[Border] = None) -> None:
        cell.font      = copy(self.font)
        cell.fill      = copy(fill if fill is not None else self.fill)
        cell.border    = copy(border if border is not None else self.border)
        cell.alignment = copy(alignment if alignment is not None else self.alignment)
        cell.number_format = number_format or self.number_format


@dataclass
class FormatProfile:
    """Everything worth cloning from the template sheet."""
    source: str
    header_row_styles: List[StyleSnapshot]
    id_style: StyleSnapshot
    delta_style: StyleSnapshot
    notes_style: StyleSnapshot
    band_fills: List[PatternFill] = field(default_factory=list)
    id_width: float = 15.0
    delta_width: float = 12.0
    notes_width: float = 28.0
    header_row_heights: List[Optional[float]] = field(default_factory=list)
    data_row_height: Optional[float] = None
    template_number_format: str = "General"
    tab_color: Optional[str] = None
    show_gridlines: bool = True

    @property
    def group_header_style(self) -> StyleSnapshot:
        return self.header_row_styles[0]

    @property
    def year_header_style(self) -> StyleSnapshot:
        return self.header_row_styles[1] if len(self.header_row_styles) > 1 else self.header_row_styles[0]

def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _fill_key(fill: PatternFill):
    """Comparable identity for a fill, so repeats can be spotted."""
    if fill is None or fill.fill_type is None:
        return ("none",)
    start = fill.start_color
    return (fill.fill_type, getattr(start, "rgb", None), getattr(start, "theme", None),
            getattr(start, "tint", None), getattr(start, "indexed", None))


def _find_first_data_row(ws, max_scan: int = 40) -> int:
    """First row holding two or more numbers. Header band is everything above it."""
    for row in range(1, min(ws.max_row, max_scan) + 1):
        numbers = sum(1 for cell in ws[row] if _is_number(cell.value))
        if numbers >= 2:
            return row
    return 2   # nothing numeric found; assume a single header row


def _numeric_columns(ws, row: int) -> List[int]:
    return [cell.column for cell in ws[row] if _is_number(cell.value)]


def _notes_column(ws, row: int, numeric_cols: List[int]) -> Optional[int]:
    """Right-most text column sitting past the numbers."""
    after = max(numeric_cols) if numeric_cols else 0
    candidates = [cell.column for cell in ws[row]
                  if cell.column > after and isinstance(cell.value, str) and cell.value.strip()]
    if candidates:
        return max(candidates)
    return after + 1 if after else None


def _representative_header_cell(ws, row: int):
    """First cell in the row carrying text, else the first cell."""
    for cell in ws[row]:
        if cell.value not in (None, ""):
            return cell
    return ws.cell(row=row, column=1)


def _width_of(ws, col: Optional[int], default: float) -> float:
    if col is None:
        return default
    dim = ws.column_dimensions.get(get_column_letter(col))
    if dim is not None and dim.width:
        return float(dim.width)
    return default

def harvest_profile(path: str, sheet_name: str) -> FormatProfile:
    """Open the template and read its formatting into a FormatProfile."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"Template workbook not found: {path}")

    wb = load_workbook(path)   # full load: read_only mode drops styling
    try:
        if sheet_name not in wb.sheetnames:
            raise KeyError(f"'{sheet_name}' not in {path}. Tabs present: {wb.sheetnames}")
        ws = wb[sheet_name]

        first_data  = _find_first_data_row(ws)
        header_rows = list(range(1, max(1, first_data)))
        if not header_rows:
            header_rows = [1]

        numeric_cols = _numeric_columns(ws, first_data)
        delta_col    = numeric_cols[0] if numeric_cols else 2
        notes_col    = _notes_column(ws, first_data, numeric_cols)
        id_col       = 1

        header_styles = [StyleSnapshot.from_cell(_representative_header_cell(ws, r))
                         for r in header_rows]

        # Row banding: distinct fills down the id column across the first dozen data rows.
        band_fills, seen = [], set()
        for row in range(first_data, min(first_data + 12, ws.max_row) + 1):
            fill = ws.cell(row=row, column=id_col).fill
            key  = _fill_key(fill)
            if key not in seen:
                seen.add(key)
                band_fills.append(copy(fill))
            if len(band_fills) >= 3:
                break

        heights = [ws.row_dimensions[r].height for r in header_rows]

        profile = FormatProfile(
            source=f"{os.path.basename(path)} :: {sheet_name}",
            header_row_styles=header_styles,
            id_style=StyleSnapshot.from_cell(ws.cell(row=first_data, column=id_col)),
            delta_style=StyleSnapshot.from_cell(ws.cell(row=first_data, column=delta_col)),
            notes_style=StyleSnapshot.from_cell(
                ws.cell(row=first_data, column=notes_col or delta_col)),
            band_fills=band_fills,
            id_width=_width_of(ws, id_col, 15.0),
            delta_width=_width_of(ws, delta_col, 12.0),
            notes_width=_width_of(ws, notes_col, 28.0),
            header_row_heights=heights,
            data_row_height=ws.row_dimensions[first_data].height,
            template_number_format=ws.cell(row=first_data, column=delta_col).number_format,
            tab_color=getattr(ws.sheet_properties.tabColor, "rgb", None),
            show_gridlines=bool(ws.sheet_view.showGridLines),
        )
        return profile
    finally:
        wb.close()


def describe_profile(profile: FormatProfile) -> None:
    """Print what was picked up, so it can be sanity-checked before the run."""
    print(f"Template : {profile.source}")
    print(f"Header rows detected : {len(profile.header_row_styles)}")
    for idx, style in enumerate(profile.header_row_styles, start=1):
        colour = getattr(style.fill.start_color, "rgb", None)
        print(f"  row {idx}: bold={style.font.b} size={style.font.sz} "
              f"font_colour={getattr(style.font.color, 'rgb', None)} fill={colour}")
    print(f"Band fills : {[getattr(f.start_color, 'rgb', None) for f in profile.band_fills] or 'none'}")
    print(f"Widths     : id={profile.id_width:.1f} delta={profile.delta_width:.1f} notes={profile.notes_width:.1f}")
    print(f"Row height : header={profile.header_row_heights} data={profile.data_row_height}")
    print(f"Number fmt : {profile.template_number_format!r}")
    print(f"Tab colour : {profile.tab_color}   gridlines={profile.show_gridlines}")

def fallback_profile() -> FormatProfile:
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centred = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header = StyleSnapshot(
        font=Font(bold=True, color="FFFFFF"),
        fill=PatternFill("solid", start_color="2F5597", end_color="2F5597"),
        border=border, alignment=centred)
    subhead = StyleSnapshot(
        font=Font(bold=True, color="FFFFFF"),
        fill=PatternFill("solid", start_color="4472C4", end_color="4472C4"),
        border=border, alignment=centred)
    body = StyleSnapshot(font=Font(), fill=PatternFill(), border=border, alignment=centred)
    notes = StyleSnapshot(font=Font(), fill=PatternFill(), border=border,
                          alignment=Alignment(horizontal="left", vertical="center", wrap_text=True))

    return FormatProfile(
        source="built-in fallback",
        header_row_styles=[header, subhead],
        id_style=body, delta_style=body, notes_style=notes,
        band_fills=[PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF"),
                    PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")],
        id_width=15.0, delta_width=12.0, notes_width=28.0,
        header_row_heights=[28, 22], data_row_height=None,
        template_number_format=DELTA_FORMAT)


def load_profile() -> FormatProfile:
    """Harvest the template; fall back with a warning if it isn't reachable."""
    try:
        profile = harvest_profile(TEMPLATE_WORKBOOK, TEMPLATE_SHEET)
        print(f"Formatting cloned from {profile.source}")
        return profile
    except Exception as exc:
        print(f"WARNING: could not read the template ({type(exc).__name__}: {exc})")
        print("         Falling back to built-in formatting. Numbers are unaffected.")
        return fallback_profile()

def _emphasise(border: Border) -> Border:
    """Same border with a heavier left edge, marking the start of a metric block."""
    return Border(left=Side(style="medium"), right=copy(border.right),
                  top=copy(border.top), bottom=copy(border.bottom))


def write_header(ws, profile: FormatProfile) -> None:
    group_style, year_style = profile.group_header_style, profile.year_header_style

    for idx, title in enumerate(ID_HEADERS, start=1):
        ws.cell(row=1, column=idx, value=title)
        ws.merge_cells(start_row=1, start_column=idx, end_row=HEADER_ROWS, end_column=idx)

    col = FIRST_DELTA_COL
    for metric in METRICS:
        ws.cell(row=1, column=col, value=metric["label"])
        ws.merge_cells(start_row=1, start_column=col,
                       end_row=1, end_column=col + len(YEAR_COLUMNS) - 1)
        for offset, label in enumerate(YEAR_LABELS):
            ws.cell(row=2, column=col + offset, value=label)
        col += len(YEAR_COLUMNS)

    for offset, title in enumerate(NOTE_HEADERS):
        idx = FIRST_NOTE_COL + offset
        ws.cell(row=1, column=idx, value=title)
        ws.merge_cells(start_row=1, start_column=idx, end_row=HEADER_ROWS, end_column=idx)

    for row in range(1, HEADER_ROWS + 1):
        style = group_style if row == 1 else year_style
        for col_idx in range(1, TOTAL_COLS + 1):
            cell = ws.cell(row=row, column=col_idx)
            border = _emphasise(style.border) if col_idx in BOUNDARY_COLS else style.border
            style.apply(cell, number_format="General", border=border)

    for offset, height in enumerate(profile.header_row_heights[:HEADER_ROWS]):
        if COPY_ROW_HEIGHTS and height:
            ws.row_dimensions[offset + 1].height = height


def style_data_cell(ws, row: int, col: int, profile: FormatProfile,
                    band: Optional[PatternFill]) -> None:
    cell = ws.cell(row=row, column=col)

    if col <= len(ID_HEADERS):
        style, number_format = profile.id_style, "General"
    elif col in METRIC_BY_COL:
        style = profile.delta_style
        number_format = (profile.template_number_format if USE_TEMPLATE_NUMBER_FORMAT
                         else METRIC_BY_COL[col]["number_format"])
    else:
        style, number_format = profile.notes_style, "General"

    border = _emphasise(style.border) if col in BOUNDARY_COLS else style.border
    style.apply(cell, number_format=number_format, fill=band, border=border)


def set_widths(ws, profile: FormatProfile) -> None:
    if not COPY_COLUMN_WIDTHS:
        return
    for col in range(1, len(ID_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = profile.id_width
    for col in range(FIRST_DELTA_COL, LAST_DELTA_COL + 1):
        ws.column_dimensions[get_column_letter(col)].width = profile.delta_width
    for col in range(FIRST_NOTE_COL, TOTAL_COLS + 1):
        ws.column_dimensions[get_column_letter(col)].width = profile.notes_width


def finish_sheet(ws, profile: FormatProfile, last_row: int) -> None:
    ws.freeze_panes = f"{get_column_letter(FIRST_DELTA_COL)}{DATA_START_ROW}"
    set_widths(ws, profile)
    if COPY_ROW_HEIGHTS and profile.data_row_height:
        for row in range(DATA_START_ROW, last_row + 1):
            ws.row_dimensions[row].height = profile.data_row_height
    if COPY_TAB_COLOR and profile.tab_color:
        ws.sheet_properties.tabColor = profile.tab_color
    if COPY_GRIDLINE_SETTING:
        ws.sheet_view.showGridLines = profile.show_gridlines

def write_flat(ws, rows: List[list], profile: FormatProfile) -> int:
    """Rows with no grouping. Banding, if any, alternates row by row."""
    write_header(ws, profile)
    for offset, row_data in enumerate(rows):
        row = DATA_START_ROW + offset
        band = profile.band_fills[offset % len(profile.band_fills)] if profile.band_fills else None
        for col, value in enumerate(row_data, start=1):
            ws.cell(row=row, column=col, value=value)
            style_data_cell(ws, row, col, profile, band)
    last_row = DATA_START_ROW + len(rows) - 1
    finish_sheet(ws, profile, last_row)
    return last_row


def write_grouped(ws, rows: List[list], profile: FormatProfile) -> int:
    """
    Rows grouped by (Product, Vantage): Product and Vantage merge vertically down
    each block, and the band shade changes per block when BAND_BY == 'group'.
    """
    write_header(ws, profile)

    grouped, order = {}, []
    for row_data in rows:
        key = (row_data[0], row_data[2])   # (Product, Vantage)
        if key not in grouped:
            grouped[key] = []
            order.append(key)
        grouped[key].append(row_data)

    current_row, group_index, row_index = DATA_START_ROW, 0, 0
    for key in order:
        block = grouped[key]

        for offset, row_data in enumerate(block):
            row = current_row + offset
            if profile.band_fills:
                pick = group_index if BAND_BY == "group" else row_index
                band = profile.band_fills[pick % len(profile.band_fills)]
            else:
                band = None
            for col, value in enumerate(row_data, start=1):
                ws.cell(row=row, column=col, value=value)
                style_data_cell(ws, row, col, profile, band)
            row_index += 1

        if len(block) > 1:
            for col in (1, 3):   # Product, Vantage
                ws.merge_cells(start_row=current_row, start_column=col,
                               end_row=current_row + len(block) - 1, end_column=col)

        current_row += len(block)
        group_index += 1

    last_row = current_row - 1
    finish_sheet(ws, profile, last_row)
    return last_row

def select_sheets(wb) -> List[str]:
    names = wb.sheetnames
    if START_SHEET in names and END_SHEET in names:
        lo, hi = sorted((names.index(START_SHEET), names.index(END_SHEET)))
        return names[lo:hi + 1]
    print("WARNING: start/end sheet not found. Using every tab except the summary.")
    return [n for n in names if n != SUMMARY_SHEET]


def build_rows(curr_wb, prev_wb):
    """Split the product tabs into those with a second dollar amount and those without."""
    with_cl, without_cl = [], []

    for sheet_name in select_sheets(curr_wb):
        first_word, product, vantage, credit_line = parse_sheet_name(sheet_name)
        prev_sheet = prev_wb[sheet_name] if sheet_name in prev_wb.sheetnames else None
        if prev_sheet is None:
            print(f"  note: '{sheet_name}' is missing from the prior workbook - deltas left blank")

        deltas = collect_deltas(curr_wb[sheet_name], prev_sheet)
        row = [product, first_word, vantage, credit_line] + deltas + [None] * len(NOTE_HEADERS)
        (with_cl if credit_line else without_cl).append(row)

    return with_cl, without_cl


def main(curr_path: str = None, prev_path: str = None,
         output_path: str = None, profile: FormatProfile = None) -> str:
    curr_path   = curr_path   or CURR_WORKBOOK
    prev_path   = prev_path   or PREV_WORKBOOK
    output_path = output_path or OUTPUT_PATH

    for label, path in (("current", curr_path), ("prior", prev_path)):
        if not os.path.exists(path):
            raise FileNotFoundError(f"Could not find the {label} workbook: {path}")

    if profile is None:
        profile = load_profile()

    print("Loading workbooks...")
    curr_wb = load_workbook(curr_path, data_only=True)
    prev_wb = load_workbook(prev_path, data_only=True)
    try:
        with_cl, without_cl = build_rows(curr_wb, prev_wb)
    finally:
        curr_wb.close()
        prev_wb.close()

    out = Workbook()
    ws_main = out.active
    ws_main.title = "With Second Dollar"
    ws_none = out.create_sheet("No Second Dollar")

    write_grouped(ws_main, with_cl, profile)
    write_flat(ws_none, without_cl, profile)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    out.save(output_path)

    print("\nDone.")
    print(f"  Metrics       : {', '.join(m['label'] for m in METRICS)}")
    print(f"  Formatting    : {profile.source}")
    print(f"  With 2nd $    : {len(with_cl)} rows")
    print(f"  Without 2nd $ : {len(without_cl)} rows")
    print(f"  Saved to      : {output_path}")
    return output_path

def run_self_test(verbose: bool = True) -> str:
    import random
    import tempfile

    tmp = tempfile.mkdtemp(prefix="finance_impact_selftest_")
    tabs = ["MoneyLion $75 601-6", "MoneyLion $75_$400 601-6",
            "Credit Karma $95_$500 <=600", "Credit Karma $95_$800 <=600",
            "Organic $95 <=600"]

    def build_forecast(path: str, offset: float) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        wb.create_sheet(SUMMARY_SHEET)
        for tab in tabs:
            ws = wb.create_sheet(tab)
            for metric in METRICS:
                for col in YEAR_COLUMNS:
                    rng = random.Random(f"{tab}{metric['row']}{col}")
                    ws[f"{col}{metric['row']}"] = round(rng.random() * 10 + offset, 4)
        wb.save(path)
        wb.close()

    def build_template(path: str) -> None:
        """Stand-in Loss Notes: merged two-row header, banded rows, notes column."""
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET

        gold = PatternFill("solid", start_color="FFC000", end_color="FFC000")
        navy = PatternFill("solid", start_color="1F3864", end_color="1F3864")
        band_a = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
        band_b = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, title in enumerate(["Product", "Yr 1 Loss", "Yr 2 Loss", "Notes"], start=1):
            top = ws.cell(row=1, column=col, value=title)
            top.font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
            top.fill = navy
            top.border = border
            top.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            sub = ws.cell(row=2, column=col, value="" if col in (1, 4) else "bps")
            sub.font = Font(bold=True, color="000000", size=10, name="Calibri")
            sub.fill = gold
            sub.border = border
            sub.alignment = Alignment(horizontal="center", vertical="center")

        for offset in range(6):
            row = 3 + offset
            fill = band_a if offset % 2 == 0 else band_b
            values = [f"Product {offset}", 1.5 + offset, 2.5 + offset, "some note"]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.fill = fill
                cell.border = border
                cell.alignment = (Alignment(horizontal="left", vertical="center", wrap_text=True)
                                  if col == 4 else
                                  Alignment(horizontal="center", vertical="center"))
                if col in (2, 3):
                    cell.number_format = "0.0"

        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 11
        ws.column_dimensions["C"].width = 11
        ws.column_dimensions["D"].width = 42
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 18
        ws.sheet_properties.tabColor = "FF0000"
        wb.save(path)
        wb.close()

    curr_path = os.path.join(tmp, "curr.xlsm")
    prev_path = os.path.join(tmp, "prev.xlsm")
    tmpl_path = os.path.join(tmp, "template.xlsx")
    out_path  = os.path.join(tmp, "finance_impact.xlsx")

    build_forecast(curr_path, 0.75)
    build_forecast(prev_path, 0.0)
    build_template(tmpl_path)

    tmpl_profile = harvest_profile(tmpl_path, TEMPLATE_SHEET)
    if verbose:
        describe_profile(tmpl_profile)
        print()

    main(curr_path=curr_path, prev_path=prev_path,
         output_path=out_path, profile=tmpl_profile)

    if verbose:
        check = load_workbook(out_path)
        ws = check["With Second Dollar"]
        print(f"\nOutput dims {ws.dimensions}, freeze {ws.freeze_panes}, tab {ws.sheet_properties.tabColor.rgb}")
        for row in ws.iter_rows(min_row=1, max_row=4, max_col=TOTAL_COLS, values_only=True):
            print(["" if v is None else v for v in row])
        probe = ws.cell(row=DATA_START_ROW, column=FIRST_DELTA_COL)
        print(f"\nDelta cell {probe.coordinate}: fill={probe.fill.start_color.rgb} "
              f"format={probe.number_format!r} font={probe.font.name}/{probe.font.sz}")
        head = ws.cell(row=1, column=FIRST_DELTA_COL)
        print(f"Header cell {head.coordinate}: fill={head.fill.start_color.rgb} "
              f"bold={head.font.b} colour={head.font.color.rgb}")
        check.close()

    return out_path


# run_self_test()

if __name__ == "__main__":
    main()
