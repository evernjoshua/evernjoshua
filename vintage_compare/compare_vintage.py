"""
Cell-by-cell value comparison of two .xlsb workbooks with a row offset.

Compares ONLY these sheets (every other sheet is never opened):
    Vintage1, Vintage2, ... Vintage24

Row alignment
-------------
A new row was inserted in the WIP file, so from row 9 down every WIP row is
one lower than in the original:
    Original row 9  <->  WIP row 10
    Original row 10 <->  WIP row 11   ... and so on.
Rows 1-8 are skipped. WIP row 9 (the new row) has no original counterpart
and is not compared.

Rules
-----
* Values only: .xlsb stores the last calculated value of every cell, and
  that is what is read. Formulas and charts are ignored.
* Column A is copied as is from the original (the row's attribute label).
  If the WIP label on the paired row is different it is flagged - a quick
  check that the rows really line up.
* Original rows 9-13 are copied as is (header block).
* Original rows 14+ show the delta:
      numbers      -> WIP - Original
      text/blank   -> blank if equal, otherwise "Original -> WIP"
* Compared columns: B..Z on Vintage1, B..Y on Vintage2..Vintage24.

Output: an .xlsx report with
    * "Summary"      - one row per sheet with the number of differences
    * "Differences"  - every differing cell with both cell addresses,
                       the attribute, both values and the delta
    * one tab per sheet laid out on the ORIGINAL row numbers, with an extra
      "WIP row" column showing the paired WIP row number

Usage
-----
    pip install pyxlsb openpyxl
    python compare_vintage.py "Digital Vantage.xlsb" "Digital by Vantage WIP.xlsb"
    python compare_vintage.py orig.xlsb wip.xlsb -o my_report.xlsx --tolerance 0.0001
"""

import argparse
import sys
from numbers import Number

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pyxlsb import open_workbook

START_ROW = 9             # first original row compared (rows 1-8 skipped)
WIP_ROW_OFFSET = 1        # WIP row = original row + 1
HEADER_LAST_ROW = 13      # original rows START_ROW..13 are printed as is
FIRST_COL = 2             # column B
FULL_LAST_COL = 26        # column Z (Vintage1)
SHORT_LAST_COL = 25       # column Y (Vintage2..Vintage24)

SHEETS = [f"Vintage{i}" for i in range(1, 25)]

DIFF_FILL = PatternFill("solid", fgColor="FFC7CE")
MISSING_FILL = PatternFill("solid", fgColor="FFEB9C")
BOLD = Font(bold=True)
GREY = Font(italic=True, color="808080")


def last_col_for(sheet_name):
    return FULL_LAST_COL if sheet_name == "Vintage1" else SHORT_LAST_COL


def read_sheet(path, sheet_name, last_col):
    """Return {(row, col): value} (1-based) for columns A..last_col, or None
    if the sheet does not exist. Only the requested sheet is read."""
    with open_workbook(path) as wb:
        if sheet_name not in wb.sheets:
            return None
        cells = {}
        with wb.get_sheet(sheet_name) as sh:
            for row in sh.rows(sparse=True):
                for c in row:
                    col = c.c + 1
                    if col > last_col or c.v is None:
                        continue
                    v = c.v
                    if isinstance(v, str) and v.strip() == "":
                        continue
                    cells[(c.r + 1, col)] = v
        return cells


def is_num(v):
    return isinstance(v, Number) and not isinstance(v, bool)


def compare_values(orig, wip, tolerance):
    """Return (is_different, delta_to_display)."""
    if is_num(orig) or is_num(wip):
        if (is_num(orig) or orig is None) and (is_num(wip) or wip is None):
            delta = (wip or 0) - (orig or 0)
            if abs(delta) <= tolerance:
                return False, 0
            return True, delta
        # number vs text
        return True, f"{orig!r} -> {wip!r}"
    if orig == wip:
        return False, None
    return True, f"{'' if orig is None else orig} -> {'' if wip is None else wip}"


def compare_sheet(sheet_name, orig_cells, wip_cells, args):
    last_col = last_col_for(sheet_name)
    offset = args.offset
    max_orig = max([r for r, _ in orig_cells] + [args.start_row])
    max_wip = max([r - offset for r, _ in wip_cells] + [args.start_row])
    max_row = max(max_orig, max_wip)

    grid = {}          # (orig_row, col) -> value to write in the report tab
    flagged = set()    # cells to highlight
    diffs = []

    for r in range(args.start_row, max_row + 1):
        wr = r + offset
        label = orig_cells.get((r, 1))
        wip_label = wip_cells.get((wr, 1))
        grid[(r, 1)] = label if label is not None else wip_label
        if label != wip_label:
            flagged.add((r, 1))
            diffs.append((sheet_name, f"A{r}", f"A{wr}", label, label, wip_label,
                          "label differs - check row alignment"))

        for c in range(FIRST_COL, last_col + 1):
            col = get_column_letter(c)
            o = orig_cells.get((r, c))
            w = wip_cells.get((wr, c))
            if r <= args.header_last_row:
                grid[(r, c)] = o if o is not None else w
                if o != w:
                    flagged.add((r, c))
                    diffs.append((sheet_name, f"{col}{r}", f"{col}{wr}", label, o, w,
                                  "header differs"))
                continue
            different, delta = compare_values(o, w, args.tolerance)
            grid[(r, c)] = delta
            if different:
                flagged.add((r, c))
                diffs.append((sheet_name, f"{col}{r}", f"{col}{wr}", label, o, w, delta))

    return grid, flagged, diffs, max_row, last_col


def write_sheet_tab(ws, grid, flagged, max_row, last_col, args):
    wip_col = last_col + 2
    ws.cell(row=args.start_row - 1, column=wip_col, value="WIP row").font = BOLD
    for r in range(args.start_row, max_row + 1):
        ws.cell(row=r, column=wip_col, value=r + args.offset).font = GREY
    for (r, c), v in grid.items():
        if v is None:
            continue
        cell = ws.cell(row=r, column=c, value=v)
        if r <= args.header_last_row or c == 1:
            cell.font = BOLD
        if (r, c) in flagged:
            cell.fill = DIFF_FILL
    ws.freeze_panes = ws.cell(row=args.header_last_row + 1, column=2)
    ws.column_dimensions["A"].width = 40
    for c in range(FIRST_COL, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("orig", nargs="?", default="Digital Vantage.xlsb")
    p.add_argument("wip", nargs="?", default="Digital by Vantage WIP.xlsb")
    p.add_argument("-o", "--output", default="vintage_comparison.xlsx")
    p.add_argument("--tolerance", type=float, default=1e-9,
                   help="numeric differences with abs value <= this are treated as equal")
    p.add_argument("--start-row", type=int, default=START_ROW,
                   help=f"first original row to compare (default {START_ROW})")
    p.add_argument("--offset", type=int, default=WIP_ROW_OFFSET,
                   help=f"WIP row = original row + offset (default {WIP_ROW_OFFSET})")
    p.add_argument("--header-last-row", type=int, default=HEADER_LAST_ROW,
                   help=f"last original row printed as is (default {HEADER_LAST_ROW})")
    args = p.parse_args()

    out = Workbook()
    summary = out.active
    summary.title = "Summary"
    summary.append(["Sheet", "Columns compared", "Rows (Original -> WIP)", "Status", "Differences"])
    details = out.create_sheet("Differences")
    details.append(["Sheet", "Original cell", "WIP cell", "Attribute (col A)",
                    "Original", "WIP", "Delta (WIP - Original)"])
    for ws in (summary, details):
        for cell in ws[1]:
            cell.font = BOLD

    total = 0
    print(f"Original: {args.orig}\nWIP     : {args.wip}")
    print(f"Original row {args.start_row} <-> WIP row {args.start_row + args.offset} "
          f"(offset +{args.offset}); rows up to {args.header_last_row} printed as is\n")
    for name in SHEETS:
        last_col = last_col_for(name)
        cols = f"B:{get_column_letter(last_col)}"
        orig_cells = read_sheet(args.orig, name, last_col)
        wip_cells = read_sheet(args.wip, name, last_col)

        if orig_cells is None or wip_cells is None:
            where = " and ".join(n for n, d in (("Original", orig_cells), ("WIP", wip_cells))
                                 if d is None)
            status = f"missing in {where}"
            summary.append([name, cols, None, status, None])
            for cell in summary[summary.max_row]:
                cell.fill = MISSING_FILL
            print(f"{name:<12} {status}")
            continue

        grid, flagged, diffs, max_row, last_col = compare_sheet(name, orig_cells, wip_cells, args)
        write_sheet_tab(out.create_sheet(name), grid, flagged, max_row, last_col, args)
        for d in diffs:
            details.append(list(d))
        total += len(diffs)
        rows = (f"{args.start_row}-{max_row} -> "
                f"{args.start_row + args.offset}-{max_row + args.offset}")
        status = "MATCH" if not diffs else "DIFFERENT"
        summary.append([name, cols, rows, status, len(diffs)])
        if diffs:
            for cell in summary[summary.max_row]:
                cell.fill = DIFF_FILL
        print(f"{name:<12} {status:<10} {len(diffs)} difference(s)")

    for ws, widths in ((summary, [14, 18, 24, 22, 14]),
                       (details, [12, 14, 10, 40, 18, 18, 26])):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    details.freeze_panes = "A2"

    out.save(args.output)
    print(f"\nTotal differences: {total}\nReport written to {args.output}")
    return 0 if total == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
