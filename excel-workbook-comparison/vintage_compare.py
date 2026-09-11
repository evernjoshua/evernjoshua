"""Delta comparison for the vintage workbooks.

Compares two structurally identical workbooks region by region and writes a third
workbook that mirrors their shape, with deltas in place of the numbers.

Values only: formulas are compared on their cached results, charts are ignored.

The point of the module is that the layout is *declared*, not discovered -- which rows
are headers, which columns are categories, where one workbook carries extra rows -- so
the alignment is explicit and auditable rather than guessed at.
"""

from __future__ import annotations

import datetime as dt
import difflib
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

__all__ = [
    "SheetSpec", "DeltaOptions", "Grid", "read_workbook", "build_row_map",
    "resolve_sheet", "compare_sheet", "compare_workbooks", "write_delta_workbook",
    "inspect_workbook", "vintage_specs",
]

XLSB_ERRORS = {
    "0x0": "#NULL!", "0x7": "#DIV/0!", "0xf": "#VALUE!", "0x17": "#REF!",
    "0x1d": "#NAME?", "0x24": "#NUM!", "0x2a": "#N/A", "0x2b": "#GETTING_DATA",
}

Grid = Dict[Tuple[int, int], Any]


# ---------------------------------------------------------------------------
#  Layout declaration
# ---------------------------------------------------------------------------

@dataclass
class SheetSpec:
    """One sheet's layout, in the row/column numbers of the BASE workbook."""

    name: str                                   # sheet name, resolved loosely against the file
    first_row: int                              # first row of the compared block
    last_row: int                               # last row of the compared block (base numbering)
    first_col: str                              # first column of the compared block, e.g. "A"
    last_col: str                               # last column, e.g. "Y"
    header_rows: Optional[Tuple[int, int]] = None   # (1, 9) -> copied through verbatim
    key_cols: Sequence[str] = ()                # category columns: copied through, checked for alignment
    ignore_cols: Sequence[str] = ()             # columns whose differences are never reported
    insertions: Sequence[Tuple[int, int]] = ()  # (after_base_row, extra rows in OTHER at that point)

    @property
    def first_col_idx(self) -> int:
        return column_index_from_string(self.first_col)

    @property
    def last_col_idx(self) -> int:
        return column_index_from_string(self.last_col)

    def col_indices(self) -> range:
        return range(self.first_col_idx, self.last_col_idx + 1)

    def key_col_indices(self) -> List[int]:
        return [column_index_from_string(c) for c in self.key_cols]

    def ignore_col_indices(self) -> List[int]:
        return [column_index_from_string(c) for c in self.ignore_cols]


@dataclass
class DeltaOptions:
    blank_as_zero: bool = True          # base 100 vs blank -> -100 (and always flagged)
    keep_matching_text: bool = True     # identical text passes through, so the sheet still reads normally
    abs_tolerance: float = 1e-9         # deltas smaller than this are written as 0
    category_similarity: float = 0.85   # difflib ratio below which category names are a MISMATCH
    zero_as_blank: bool = False         # write blank instead of 0 where nothing moved
    max_issues: int = 50_000


def vintage_specs(n_vintages: int = 24,
                  summary_name: str = "Vintage Summary",
                  vintage_name: str = "Vintage{}") -> List[SheetSpec]:
    """The layout as specified: one summary sheet plus n vintage sheets."""
    specs = [SheetSpec(
        name=summary_name,
        header_rows=(1, 24),            # the charts block, copied through, never compared
        first_row=25, last_row=756,
        first_col="A", last_col="Y",
        key_cols=("A",),
        insertions=(),                  # no row shift on the summary
    )]
    specs += [SheetSpec(
        name=vintage_name.format(i),
        header_rows=(1, 9),
        first_row=10, last_row=248,
        first_col="A", last_col="AA",
        key_cols=("A", "B"),
        ignore_cols=("A",),             # column A differences are ignored entirely
        insertions=((200, 1),),         # OTHER has one extra row after base row 200
    ) for i in range(1, n_vintages + 1)]
    return specs


def build_row_map(spec: SheetSpec) -> Tuple[Dict[int, int], List[int]]:
    """base row -> other row, plus the OTHER rows skipped by the insertions."""
    insertions = dict(spec.insertions)
    mapping: Dict[int, int] = {}
    skipped: List[int] = []
    offset = 0
    for r in range(spec.first_row, spec.last_row + 1):
        mapping[r] = r + offset
        extra = insertions.get(r, 0)
        for k in range(1, extra + 1):
            skipped.append(r + offset + k)
        offset += extra
    return mapping, skipped


# ---------------------------------------------------------------------------
#  Reading: .xlsx / .xlsm via openpyxl, .xlsb via pyxlsb, .xls via xlrd
# ---------------------------------------------------------------------------

def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def read_workbook(path: str | Path, sheets: Optional[Iterable[str]] = None) -> Dict[str, Grid]:
    """Read cached values into {sheet name: {(row, col): value}}, 1-indexed like Excel."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Workbook not found: {p.resolve()}")
    suffix = p.suffix.lower()
    wanted = set(sheets) if sheets is not None else None

    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _read_openpyxl(p, wanted)
    if suffix == ".xlsb":
        return _read_pyxlsb(p, wanted)
    if suffix == ".xls":
        return _read_xlrd(p, wanted)
    raise ValueError(f"{p.name}: unsupported extension '{suffix}'")


def _read_openpyxl(p: Path, wanted: Optional[set]) -> Dict[str, Grid]:
    wb = load_workbook(p, data_only=True, read_only=True, keep_links=False)
    try:
        out: Dict[str, Grid] = {}
        for name in wb.sheetnames:
            if wanted is not None and name not in wanted:
                continue
            ws = wb[name]
            if not hasattr(ws, "iter_rows"):        # chart sheet
                continue
            grid: Grid = {}
            for row in ws.iter_rows():
                for cell in row:
                    if not _blank(cell.value):
                        grid[(cell.row, cell.column)] = cell.value
            out[name] = grid
        return out
    finally:
        wb.close()


def _read_pyxlsb(p: Path, wanted: Optional[set]) -> Dict[str, Grid]:
    import pyxlsb                                   # pip install pyxlsb

    out: Dict[str, Grid] = {}
    with pyxlsb.open_workbook(str(p)) as wb:
        for name in wb.sheets:
            if wanted is not None and name not in wanted:
                continue
            grid: Grid = {}
            with wb.get_sheet(name) as sh:
                for row in sh.rows(sparse=True):
                    for cell in row:
                        v = cell.v
                        if _blank(v):
                            continue
                        if isinstance(v, str) and v in XLSB_ERRORS:
                            v = XLSB_ERRORS[v]
                        grid[(cell.r + 1, cell.c + 1)] = v     # pyxlsb is 0-indexed
            out[name] = grid
    return out


def _read_xlrd(p: Path, wanted: Optional[set]) -> Dict[str, Grid]:
    import xlrd                                      # pip install xlrd

    book = xlrd.open_workbook(str(p), on_demand=True)
    try:
        out: Dict[str, Grid] = {}
        for name in book.sheet_names():
            if wanted is not None and name not in wanted:
                continue
            sh = book.sheet_by_name(name)
            grid: Grid = {}
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    ct, v = sh.cell_type(r, c), sh.cell_value(r, c)
                    if ct in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or _blank(v):
                        continue
                    if ct == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    elif ct == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    elif ct == xlrd.XL_CELL_ERROR:
                        v = xlrd.error_text_from_code.get(v, f"#ERR:{v}")
                    grid[(r + 1, c + 1)] = v
            out[name] = grid
            book.unload_sheet(name)
        return out
    finally:
        book.release_resources()


# ---------------------------------------------------------------------------
#  Sheet-name resolution: real files rarely spell sheet names the way a spec does
# ---------------------------------------------------------------------------

def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name).lower())


def resolve_sheet(wanted: str, available: Sequence[str]) -> Optional[str]:
    """Match 'Vintage1' against 'Vintage 1', 'VINTAGE_1', 'vintage1 ' and so on."""
    if wanted in available:
        return wanted
    target = _slug(wanted)
    for name in available:
        if _slug(name) == target:
            return name
    m = re.fullmatch(r"([a-z]+?)0*(\d+)", target)     # vintage01 == vintage1
    if m:
        stem, num = m.group(1), int(m.group(2))
        for name in available:
            m2 = re.fullmatch(r"([a-z]+?)0*(\d+)", _slug(name))
            if m2 and m2.group(1) == stem and int(m2.group(2)) == num:
                return name
    return None


# ---------------------------------------------------------------------------
#  Value semantics
# ---------------------------------------------------------------------------

def is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and not (
        isinstance(v, float) and math.isnan(v))


def norm_text(v: Any) -> str:
    s = re.sub(r"\s+", " ", str(v)).strip().casefold()
    return re.sub(r"[^\w %&/().-]+", "", s)


def category_status(base_v: Any, other_v: Any, opt: DeltaOptions) -> Tuple[str, float]:
    """('match' | 'similar' | 'mismatch' | 'blank', similarity ratio)."""
    if _blank(base_v) and _blank(other_v):
        return "match", 1.0
    if _blank(base_v) or _blank(other_v):
        return "blank", 0.0
    a, b = norm_text(base_v), norm_text(other_v)
    if a == b:
        return "match", 1.0
    ratio = difflib.SequenceMatcher(None, a, b).ratio()
    return ("similar" if ratio >= opt.category_similarity else "mismatch"), ratio


def cell_delta(base_v: Any, other_v: Any, opt: DeltaOptions) -> Tuple[Any, Optional[str]]:
    """Return (value to write, issue label or None)."""
    b_blank, o_blank = _blank(base_v), _blank(other_v)

    if b_blank and o_blank:
        return None, None

    if is_number(base_v) and is_number(other_v):
        d = float(other_v) - float(base_v)
        if abs(d) <= opt.abs_tolerance:
            return (None if opt.zero_as_blank else 0), None
        return d, None

    if b_blank or o_blank:
        present, missing_side = (other_v, "base") if b_blank else (base_v, "other")
        if is_number(present):
            if opt.blank_as_zero:
                d = float(other_v or 0) - float(base_v or 0)
                return d, f"blank in {missing_side} workbook, treated as 0"
            return None, f"blank in {missing_side} workbook, not compared"
        return None, f"blank in {missing_side} workbook"

    if not is_number(base_v) and not is_number(other_v):
        if norm_text(base_v) == norm_text(other_v):
            return (base_v if opt.keep_matching_text else None), None
        return f"{base_v} -> {other_v}", "text differs"

    return f"{base_v} -> {other_v}", "number vs text"


# ---------------------------------------------------------------------------
#  Comparison
# ---------------------------------------------------------------------------

@dataclass
class SheetResult:
    spec: SheetSpec
    base_sheet: str
    other_sheet: str
    values: Grid = field(default_factory=dict)      # what to write into the output sheet
    issues: List[Dict[str, Any]] = field(default_factory=list)
    skipped_other_rows: List[int] = field(default_factory=list)
    cells_compared: int = 0
    nonzero_deltas: int = 0
    max_abs_delta: float = 0.0
    max_abs_cell: str = ""
    category_mismatches: int = 0


def compare_sheet(spec: SheetSpec, base: Grid, other: Grid,
                  base_sheet: str, other_sheet: str, opt: DeltaOptions) -> SheetResult:
    res = SheetResult(spec=spec, base_sheet=base_sheet, other_sheet=other_sheet)
    row_map, res.skipped_other_rows = build_row_map(spec)
    key_idx = set(spec.key_col_indices())
    ignore_idx = set(spec.ignore_col_indices())

    if spec.header_rows:                                    # copied through verbatim from base
        h_first, h_last = spec.header_rows
        for (r, c), v in base.items():
            if h_first <= r <= h_last:
                res.values[(r, c)] = v

    for base_row, other_row in row_map.items():
        for col in spec.col_indices():
            bv = base.get((base_row, col))
            ov = other.get((other_row, col))
            ref = f"{get_column_letter(col)}{base_row}"

            if col in key_idx:
                res.values[(base_row, col)] = bv         # categories always come from base
                if col in ignore_idx:
                    continue
                status, ratio = category_status(bv, ov, opt)
                if status in {"mismatch", "blank"} and not (_blank(bv) and _blank(ov)):
                    res.category_mismatches += 1
                    res.issues.append({
                        "Sheet": base_sheet, "Cell": ref, "Base row": base_row,
                        "Other row": other_row, "Column": get_column_letter(col),
                        "Issue": f"category {status} (similarity {ratio:.2f})",
                        "Base value": bv, "Other value": ov,
                    })
                continue

            if col in ignore_idx:
                res.values[(base_row, col)] = bv
                continue

            value, issue = cell_delta(bv, ov, opt)
            res.values[(base_row, col)] = value
            res.cells_compared += 1
            if is_number(value) and value != 0:
                res.nonzero_deltas += 1
                if abs(value) > res.max_abs_delta:
                    res.max_abs_delta, res.max_abs_cell = abs(value), ref
            if issue and len(res.issues) < opt.max_issues:
                res.issues.append({
                    "Sheet": base_sheet, "Cell": ref, "Base row": base_row,
                    "Other row": other_row, "Column": get_column_letter(col),
                    "Issue": issue, "Base value": bv, "Other value": ov,
                })
    return res


def compare_workbooks(base_path, other_path, specs: Sequence[SheetSpec],
                      opt: Optional[DeltaOptions] = None) -> Dict[str, Any]:
    opt = opt or DeltaOptions()
    wanted = [s.name for s in specs]

    base_all = read_workbook(base_path)
    other_all = read_workbook(other_path)
    base_names, other_names = list(base_all), list(other_all)

    results: List[SheetResult] = []
    unresolved: List[Dict[str, str]] = []
    for spec in specs:
        b_name = resolve_sheet(spec.name, base_names)
        o_name = resolve_sheet(spec.name, other_names)
        if b_name is None or o_name is None:
            unresolved.append({
                "Sheet": spec.name,
                "In base": b_name or "NOT FOUND",
                "In other": o_name or "NOT FOUND",
            })
            continue
        results.append(compare_sheet(spec, base_all[b_name], other_all[o_name],
                                     b_name, o_name, opt))

    return {
        "base_path": str(Path(base_path).resolve()),
        "other_path": str(Path(other_path).resolve()),
        "base_sheets": base_names,
        "other_sheets": other_names,
        "results": results,
        "unresolved": unresolved,
        "options": opt,
        "generated": dt.datetime.now(),
    }


# ---------------------------------------------------------------------------
#  Inspection: check the declared layout against what the files actually contain
# ---------------------------------------------------------------------------

def inspect_workbook(path, specs: Sequence[SheetSpec]) -> List[Dict[str, Any]]:
    """Per spec: does the sheet exist, and where does its data actually start and stop?"""
    all_grids = read_workbook(path)
    names = list(all_grids)
    rows = []
    for spec in specs:
        resolved = resolve_sheet(spec.name, names)
        if resolved is None:
            rows.append({"Spec sheet": spec.name, "Found as": "NOT FOUND"})
            continue
        grid = all_grids[resolved]
        rs = [r for r, _ in grid] or [0]
        cs = [c for _, c in grid] or [0]
        in_block = [r for r in rs if spec.first_row <= r <= spec.last_row]
        rows.append({
            "Spec sheet": spec.name,
            "Found as": resolved,
            "Populated rows": f"{min(rs)}-{max(rs)}",
            "Populated cols": f"{get_column_letter(min(cs))}-{get_column_letter(max(cs))}",
            "Spec block": f"{spec.first_col}{spec.first_row}:{spec.last_col}{spec.last_row}",
            "Rows in block": len(set(in_block)),
            "Data past last_row": max(rs) - spec.last_row if max(rs) > spec.last_row else 0,
            "Cells": len(grid),
        })
    return rows


# ---------------------------------------------------------------------------
#  Output workbook
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_MOVED_FILL = PatternFill("solid", fgColor="FFF2CC")

ISSUE_COLUMNS = ["Sheet", "Cell", "Base row", "Other row", "Column",
                 "Issue", "Base value", "Other value"]


def _safe(v: Any, limit: int = 500) -> Any:
    if v is None or isinstance(v, (int, float, bool, dt.datetime, dt.date, dt.time)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return str(v)
        return v
    s = ILLEGAL_CHARACTERS_RE.sub("", v if isinstance(v, str) else repr(v))
    return s if len(s) <= limit else s[:limit] + f"... [{len(s):,} chars]"


def _put(ws, row: int, col: int, value: Any):
    v = _safe(value)
    cell = ws.cell(row=row, column=col, value=v)
    if isinstance(v, str) and v[:1] in {"=", "+", "-", "@"}:
        cell.data_type = "s"            # never let a value become a formula
    return cell


def _table(ws, headers: Sequence[str], rows: Sequence[Sequence[Any]], start_row: int = 1):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=i, value=h)
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
    for j, r in enumerate(rows, start=start_row + 1):
        for i, v in enumerate(r, start=1):
            _put(ws, j, i, v)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    if rows:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def write_delta_workbook(result: Dict[str, Any], out_path: str | Path) -> Path:
    """Mirror the input layout, with deltas as the cell values, plus Summary and Issues tabs."""
    opt: DeltaOptions = result["options"]
    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet("_Summary")
    issues_ws = wb.create_sheet("_Issues")

    for res in result["results"]:
        ws = wb.create_sheet(res.base_sheet[:31])
        for (r, c), v in res.values.items():
            _put(ws, r, c, v)

        spec = res.spec
        first_delta_col = spec.first_col_idx
        key_and_ignored = set(spec.key_col_indices()) | set(spec.ignore_col_indices())
        while first_delta_col in key_and_ignored and first_delta_col <= spec.last_col_idx:
            first_delta_col += 1
        if first_delta_col <= spec.last_col_idx:
            rng = (f"{get_column_letter(first_delta_col)}{spec.first_row}:"
                   f"{spec.last_col}{spec.last_row}")
            ws.conditional_formatting.add(rng, CellIsRule(
                operator="notEqual", formula=["0"], fill=_MOVED_FILL))
        ws.freeze_panes = ws.cell(row=spec.first_row, column=first_delta_col)

    # ---- Summary -----------------------------------------------------------
    summary["A1"] = "Vintage delta comparison"
    summary["A1"].font = Font(bold=True, size=14, color="1F3864")
    meta = [
        ("Base workbook (A)", result["base_path"]),
        ("Other workbook (B)", result["other_path"]),
        ("Deltas are", "B minus A"),
        ("Generated", result["generated"].strftime("%Y-%m-%d %H:%M:%S")),
        ("Blank treated as zero", opt.blank_as_zero),
        ("Category similarity threshold", opt.category_similarity),
        ("Numeric tolerance", opt.abs_tolerance),
    ]
    row = 3
    for label, value in meta:
        summary.cell(row=row, column=1, value=label).font = Font(bold=True)
        _put(summary, row, 2, value)
        row += 1

    headers = ["Sheet", "Base sheet", "Other sheet", "Block", "Rows compared",
               "Row shift", "Skipped rows in B", "Cells compared", "Non-zero deltas",
               "Largest |delta|", "At cell", "Category mismatches"]
    rows = []
    for res in result["results"]:
        spec = res.spec
        row_map, _ = build_row_map(spec)
        shifts = sorted({o - b for b, o in row_map.items()})
        rows.append([
            spec.name, res.base_sheet, res.other_sheet,
            f"{spec.first_col}{spec.first_row}:{spec.last_col}{spec.last_row}",
            len(row_map),
            "none" if shifts == [0] else f"+{max(shifts)} after row {spec.insertions[0][0]}"
            if spec.insertions else f"{shifts}",
            ", ".join(str(r) for r in res.skipped_other_rows) or "none",
            res.cells_compared, res.nonzero_deltas,
            round(res.max_abs_delta, 6) if res.max_abs_delta else 0,
            res.max_abs_cell or "", res.category_mismatches,
        ])
    _table(summary, headers, rows, start_row=row + 1)

    if result["unresolved"]:
        start = row + len(rows) + 4
        summary.cell(row=start, column=1, value="Sheets that could not be matched").font = \
            Font(bold=True, color="C00000")
        _table(summary, ["Sheet", "In base", "In other"],
               [[u["Sheet"], u["In base"], u["In other"]] for u in result["unresolved"]],
               start_row=start + 1)

    # ---- Issues ------------------------------------------------------------
    all_issues = [i for res in result["results"] for i in res.issues]
    _table(issues_ws, ISSUE_COLUMNS,
           [[i[h] for h in ISSUE_COLUMNS] for i in all_issues])
    if not all_issues:
        issues_ws["A2"] = "No category mismatches or blank/type problems found."

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out.resolve()


def print_summary(result: Dict[str, Any], out_path: Optional[Path] = None) -> None:
    bar = "=" * 78
    print(bar)
    print(f"A (base) : {result['base_path']}")
    print(f"B (other): {result['other_path']}")
    print(f"deltas   : B minus A")
    print(bar)
    print(f"{'sheet':<20}{'rows':>6}{'cells':>9}{'moved':>9}{'largest |delta|':>18}{'cat?':>7}")
    for res in result["results"]:
        rows, _ = build_row_map(res.spec)
        print(f"{res.spec.name:<20}{len(rows):>6}{res.cells_compared:>9,}"
              f"{res.nonzero_deltas:>9,}{res.max_abs_delta:>18,.2f}"
              f"{res.category_mismatches:>7}")
    total_moved = sum(r.nonzero_deltas for r in result["results"])
    total_cat = sum(r.category_mismatches for r in result["results"])
    print(bar)
    print(f"{total_moved:,} cells moved; {total_cat:,} category mismatches")
    for u in result["unresolved"]:
        print(f"UNMATCHED SHEET: {u['Sheet']} (base={u['In base']}, other={u['In other']})")
    if out_path:
        print(f"Written to {out_path}")
